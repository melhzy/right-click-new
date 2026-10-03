#!/usr/bin/env python3
"""Build every template and check it with independent readers.

Needs: python-docx openpyxl python-pptx pillow nbformat
Optional: LibreOffice (soffice) - if present, each Office file is also
converted to PDF as an "does a real office suite open it" check.
"""
import io
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.dom import minidom
from xml.etree import ElementTree

import docx
import nbformat
import openpyxl
import pptx
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
failures = 0


def check(name, fn):
    global failures
    try:
        print(f"PASS  {name}: {fn()}")
    except Exception as exc:  # noqa: BLE001
        failures += 1
        print(f"FAIL  {name}: {type(exc).__name__}: {exc}")


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="right-click-new-"))
    out = tmp / "templates"
    subprocess.run([sys.executable, str(ROOT / "build_templates.py"), "build", str(out)],
                   check=True, stdout=subprocess.DEVNULL)
    office = {
        "docx": out / "Microsoft Word Document.docx",
        "xlsx": out / "Microsoft Excel Worksheet.xlsx",
        "pptx": out / "Microsoft PowerPoint Presentation.pptx",
    }

    def wellformed(path):
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            assert names[0] == "[Content_Types].xml", names[0]
            for n in names:
                minidom.parseString(z.read(n))
        return f"{len(names)} parts, XML well-formed"

    for kind, path in office.items():
        check(f"{kind} package", lambda p=path: wellformed(p))

    def t_docx():
        d = docx.Document(office["docx"])
        d.add_paragraph("hello")
        d.save(tmp / "rt.docx")
        s = d.sections[0]
        return f"{s.page_width.inches:.2f}x{s.page_height.inches:.2f} in, round-trips"

    def t_xlsx():
        wb = openpyxl.load_workbook(office["xlsx"])
        wb.active["A1"] = 42
        wb.save(tmp / "rt.xlsx")
        return f"sheets={wb.sheetnames}, round-trips"

    def t_pptx():
        p = pptx.Presentation(office["pptx"])
        p.slides[0].shapes.title.text = "Title"
        p.slides.add_slide(p.slide_layouts[1]).shapes.title.text = "Second"
        p.save(tmp / "rt.pptx")
        return f"layouts={[l.name for l in p.slide_layouts]}, round-trips"

    def t_ora():
        path = out / "Pinta Image.ora"
        raw = path.read_bytes()
        # The freedesktop MIME magic for image/openraster:
        assert raw[30:38] == b"mimetype" and raw[38:54] == b"image/openraster"
        with zipfile.ZipFile(path) as z:
            first = z.infolist()[0]
            assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED
            root = ElementTree.fromstring(z.read("stack.xml"))
            w, h = int(root.get("w")), int(root.get("h"))
            layers = root.findall("./stack/layer")
            assert len(layers) == 1 and layers[0].get("name") == "Background"
            layer = Image.open(io.BytesIO(z.read(layers[0].get("src"))))
            merged = Image.open(io.BytesIO(z.read("mergedimage.png")))
            thumb = Image.open(io.BytesIO(z.read("Thumbnails/thumbnail.png")))
            assert layer.size == merged.size == (w, h)
            assert layer.getpixel((0, 0)) == (255, 255, 255, 255)
            assert max(thumb.size) <= 256
        detail = f"{w}x{h}, 1 layer, thumbnail {thumb.size}"
        try:
            from pyora import Project
        except ImportError:
            return detail + " (pyora not installed)"
        project = Project.load(str(path))
        assert tuple(project.dimensions) == (w, h)
        return detail + ", opens in pyora"

    def t_sqlite():
        path = out / "SQLite Database.sqlite3"
        assert path.read_bytes()[:16] == b"SQLite format 3\x00"
        con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
        return f"integrity={con.execute('pragma integrity_check').fetchone()[0]}"

    def t_ipynb():
        nb = nbformat.read(out / "Jupyter Notebook.ipynb", as_version=4)
        nbformat.validate(nb)
        return f"nbformat {nb.nbformat}.{nb.nbformat_minor}"

    def t_python():
        compile((out / "Python Script.py").read_text(), "Python Script.py", "exec")
        return "compiles"

    def t_deterministic():
        again = tmp / "again"
        subprocess.run([sys.executable, str(ROOT / "build_templates.py"), "build", str(again)],
                       check=True, stdout=subprocess.DEVNULL)
        for p in office.values():
            assert p.read_bytes() == (again / p.name).read_bytes(), p.name
        return "Office files are byte-identical across builds"

    check("docx (python-docx)", t_docx)
    check("xlsx (openpyxl)", t_xlsx)
    check("pptx (python-pptx)", t_pptx)
    check("Pinta Image (OpenRaster)", t_ora)
    check("sqlite3", t_sqlite)
    check("ipynb (nbformat)", t_ipynb)
    check("python script", t_python)
    check("determinism", t_deterministic)

    soffice = shutil.which("soffice")
    if soffice:
        def t_libreoffice():
            lo = tmp / "lo"
            lo.mkdir()
            for kind, p in office.items():
                shutil.copy(p, lo / f"t.{kind}")
                subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                                "--outdir", str(lo / kind), str(lo / f"t.{kind}")],
                               check=True, capture_output=True, timeout=180)
                assert (lo / kind / "t.pdf").stat().st_size > 0, kind
            return "docx/xlsx/pptx convert to PDF"
        check("LibreOffice", t_libreoffice)
    else:
        print("SKIP  LibreOffice not installed")

    shutil.rmtree(tmp, ignore_errors=True)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
